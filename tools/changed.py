"""List the published files that carry a real change.

Used by .github/workflows/refresh.yml after it stages docs/data and
docs/downloads. Prints each changed file, one per line, and nothing when the
run changed no published value, so the refresh commits only when a source or
the code did.

Some fields move on every run: each source's retrieval date, the geocoder's
period, which is the build date, the build date in status.json, and the
timestamps inside the zip and the workbook. They are set aside before
comparing, and publish along with any real change. A stale-source warning
that appears or clears is a real change.

    python tools/changed.py                     staged files against HEAD
    python tools/changed.py --base A --head B   one commit against another
"""

import argparse
import csv
import io
import json
import re
import subprocess
import zipfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
PUBLISHED = ["docs/data", "docs/downloads"]

# Sources whose period is the build date rather than a period of data.
BUILD_DATED = {"geosearch"}
SOURCES_CSV = "sources.csv"
SOURCES_SHEET = "Sources and Coverage"


def git(*args):
    return subprocess.run(["git", *args], cwd=ROOT, check=True,
                          capture_output=True).stdout


def read(rev, path):
    """A file's bytes at a commit, or in the index when rev is None."""
    try:
        return git("show", f"{rev or ''}:{path}")
    except subprocess.CalledProcessError:
        return None


def steady_sources(rows):
    """Source rows without the fields that carry the date of the run."""
    out = []
    for row in rows:
        row = {k: v for k, v in row.items() if k != "retrieved"}
        if row.get("source_id") in BUILD_DATED:
            row.pop("earliest_period", None)
            row.pop("latest_period", None)
        out.append(row)
    return out


def steady_zip(data):
    archive = zipfile.ZipFile(io.BytesIO(data))
    out = {}
    for name in archive.namelist():
        content = archive.read(name)
        if name == "README.txt":
            content = re.sub(rb"^Generated \d{4}-\d{2}-\d{2}", b"Generated", content,
                             flags=re.M)
        elif name == SOURCES_CSV:
            content = steady_sources(csv.DictReader(io.StringIO(content.decode())))
        out[name] = content
    return out


def steady_workbook(data):
    """Cell values per sheet, plus every other part of the file but its timestamps.

    Sheets compare by value, since a new retrieval date can renumber the
    workbook's shared strings without any figure changing.
    """
    archive = zipfile.ZipFile(io.BytesIO(data))
    out = {}
    for name in archive.namelist():
        if name == "xl/sharedStrings.xml" or name.startswith("xl/worksheets/"):
            continue
        content = archive.read(name)
        if name == "docProps/core.xml":
            content = re.sub(rb"<dcterms:(created|modified)[^>]*>[^<]*</dcterms:\1>",
                             b"", content)
        out[name] = content
    book = load_workbook(io.BytesIO(data), read_only=True)
    for sheet in book.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if sheet.title == SOURCES_SHEET and rows:
            rows = steady_sources(dict(zip(rows[0], r)) for r in rows[1:])
        out[f"sheet:{sheet.title}"] = rows
    return out


def steady(path, data):
    """What counts as the content of a published file."""
    if path == "docs/data/status.json":
        return sorted(s["source_id"] for s in json.loads(data).get("stale_sources") or [])
    if path == "docs/data/sources.json":
        return steady_sources(json.loads(data))
    if path.endswith(".zip"):
        return steady_zip(data)
    if path.endswith(".xlsx"):
        return steady_workbook(data)
    return data


def changed(base, head):
    diff = (["diff", "--name-only", "--no-renames", base, head] if head
            else ["diff", "--cached", "--name-only", "--no-renames", base])
    paths = git(*diff, "--", *PUBLISHED).decode().split()
    found = []
    for path in paths:
        old, new = read(base, path), read(head, path)
        if old is None or new is None or steady(path, old) != steady(path, new):
            found.append(path)
    return found


def main():
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--base", default="HEAD", help="commit to compare against")
    parser.add_argument("--head", help="commit to compare; the index when left out")
    args = parser.parse_args()
    for path in changed(args.base, args.head):
        print(path)


if __name__ == "__main__":
    main()
